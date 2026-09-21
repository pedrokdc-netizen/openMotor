"""Generate a bilingual engineering report from an openMotor motor file."""

import argparse
import json
import math
import sys
import tempfile
import warnings
from datetime import date
from pathlib import Path

import numpy as np
import yaml
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from matplotlib.figure import Figure

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from motorlib.motor import Motor
from motorlib.simResult import SimAlertLevel
from tools.openmotor_optimizer import load_motor


DEFAULT_TEMPLATE = Path(__file__).with_name("report-template.yaml")


def load_report_config(path):
    configPath = Path(path).resolve()
    config = yaml.safe_load(configPath.read_text(encoding="utf-8"))
    if not isinstance(config, dict):
        raise ValueError("Report configuration must be a YAML mapping")
    for field in ("motor", "output"):
        if not config.get(field):
            raise ValueError(f"Missing required report field: {field}")

    for field in ("motor", "output", "template"):
        value = config.get(field)
        if value:
            resolved = Path(value)
            if not resolved.is_absolute():
                resolved = configPath.parent / resolved
            config[field] = str(resolved.resolve())
    config.setdefault("template", str(DEFAULT_TEMPLATE.resolve()))
    config.setdefault("metadata", {})
    config.setdefault("simulation", {})
    config.setdefault("figures", {})
    return config


def validate_report_config(config):
    if not Path(config["motor"]).is_file():
        raise ValueError(f"Motor file does not exist: {config['motor']}")
    if not Path(config["template"]).is_file():
        raise ValueError(f"Report template does not exist: {config['template']}")
    if Path(config["output"]).suffix.lower() != ".docx":
        raise ValueError("Report output must have the .docx extension")


def load_template(path):
    template = yaml.safe_load(Path(path).read_text(encoding="utf-8"))
    for field in ("document", "sections", "text", "styles"):
        if field not in template:
            raise ValueError(f"Template is missing section: {field}")
    return template


def simulate(config):
    motorData, _ = load_motor(config["motor"])
    motorData = json.loads(json.dumps(motorData))
    mapDim = config.get("simulation", {}).get("map_dim")
    if mapDim:
        motorData["config"]["mapDim"] = int(mapDim)
    motor = Motor(motorData)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        result = motor.runSimulation()
    if not result.success or not result.channels["time"].data:
        errors = [
            alert.description
            for alert in result.alerts
            if alert.level == SimAlertLevel.ERROR
        ]
        raise RuntimeError(
            "Simulation failed: " + ("; ".join(errors) or "unknown error")
        )
    return motor, result


def performance_metrics(result):
    peakPressure = result.getMaxPressure()
    averagePressure = result.getAveragePressure()
    return [
        ("Motor designation / Designacao", result.getDesignation(), ""),
        ("Total impulse / Impulso total", result.getImpulse(), "N s"),
        ("Burn time / Tempo de queima", result.getBurnTime(), "s"),
        ("Average thrust / Empuxo medio", result.getAverageForce(), "N"),
        (
            "Peak thrust / Empuxo maximo",
            result.channels["force"].getMax(),
            "N",
        ),
        (
            "Peak chamber pressure / Pressao maxima",
            peakPressure / 1e6,
            "MPa",
        ),
        (
            "Average chamber pressure / Pressao media",
            averagePressure / 1e6,
            "MPa",
        ),
        (
            "Peak - average pressure / Pressao maxima - media",
            (peakPressure - averagePressure) / 1e6,
            "MPa",
        ),
        (
            "Relative pressure spread / Variacao relativa de pressao",
            (peakPressure - averagePressure) / averagePressure,
            "",
        ),
        (
            "Propellant mass / Massa de propelente",
            result.getPropellantMass(),
            "kg",
        ),
        ("Specific impulse / Impulso especifico", result.getISP(), "s"),
        ("Peak Kn / Kn maximo", result.getPeakKN(), ""),
        (
            "Peak mass flux / Fluxo de massa maximo",
            result.getPeakMassFlux(),
            "kg/(m^2 s)",
        ),
        (
            "Peak core Mach / Mach maximo no canal",
            result.getPeakMachNumber(),
            "",
        ),
        (
            "Initial port/throat ratio / Razao canal/garganta inicial",
            result.getPortRatio(),
            "",
        ),
    ]


def format_value(value, unit=""):
    if value is None:
        return "N/A"
    if isinstance(value, str):
        return value
    if not math.isfinite(float(value)):
        return "N/A"
    absolute = abs(float(value))
    if absolute >= 1000:
        text = f"{value:,.1f}"
    elif absolute >= 10:
        text = f"{value:.3f}"
    else:
        text = f"{value:.5g}"
    return f"{text} {unit}".strip()


def property_value(name, value):
    lower = name.lower()
    if not isinstance(value, (int, float)):
        return str(value)
    if "pressure" in lower:
        return format_value(value / 1e6, "MPa")
    if any(
        token in lower
        for token in ("diameter", "length", "throat", "exit")
    ):
        return format_value(value * 1000, "mm")
    if "angle" in lower:
        return format_value(value, "deg")
    if lower == "density":
        return format_value(value, "kg/m^3")
    if "timestep" in lower:
        return format_value(value, "s")
    return format_value(value)


def generate_performance_figure(result, path, dpi):
    time = np.asarray(result.channels["time"].getData())
    pressure = np.asarray(result.channels["pressure"].getData()) / 1e6
    thrust = np.asarray(result.channels["force"].getData())
    figure = Figure(figsize=(8.2, 7.0))
    pressureAxis, thrustAxis = figure.subplots(2, 1, sharex=True)
    pressureAxis.plot(time, pressure, color="#17365D", linewidth=1.6)
    pressureAxis.set_ylabel("Pressure / Pressao (MPa)")
    pressureAxis.grid(True, alpha=0.25)
    thrustAxis.plot(time, thrust, color="#C65911", linewidth=1.6)
    thrustAxis.set_xlabel("Time / Tempo (s)")
    thrustAxis.set_ylabel("Thrust / Empuxo (N)")
    thrustAxis.grid(True, alpha=0.25)
    figure.tight_layout()
    figure.savefig(path, dpi=dpi)


def generate_diagnostics_figure(result, path, dpi):
    time = np.asarray(result.channels["time"].getData())
    kn = np.asarray(result.channels["kn"].getData())
    regression = np.asarray(result.channels["regression"].getData())
    massFlux = np.asarray(result.channels["massFlux"].getData())
    mass = np.asarray(result.channels["mass"].getData())
    figure = Figure(figsize=(8.2, 7.0))
    axes = figure.subplots(2, 2)
    axes[0, 0].plot(time, kn)
    axes[0, 0].set_ylabel("Kn")
    axes[0, 1].plot(time, np.sum(mass, axis=1))
    axes[0, 1].set_ylabel("Propellant mass / Massa (kg)")
    if regression.ndim == 2:
        for index in range(regression.shape[1]):
            axes[1, 0].plot(time, regression[:, index] * 1000)
    axes[1, 0].set_ylabel("Regression / Regressao (mm)")
    if massFlux.ndim == 2:
        axes[1, 1].plot(time, np.max(massFlux, axis=1))
    axes[1, 1].set_ylabel("Peak mass flux / Fluxo max.")
    for axis in axes.flat:
        axis.set_xlabel("Time / Tempo (s)")
        axis.grid(True, alpha=0.25)
    figure.tight_layout()
    figure.savefig(path, dpi=dpi)


def generate_regression_figure(motor, path, dpi, mapDim, contourCount):
    columns = min(3, len(motor.grains))
    rows = math.ceil(len(motor.grains) / columns)
    figure = Figure(figsize=(8.2, max(3.0, rows * 2.8)))
    axes = np.atleast_1d(
        figure.subplots(rows, columns, squeeze=False)
    ).reshape(-1)
    for index, (axis, grain) in enumerate(zip(axes, motor.grains)):
        face, _, contours, _ = grain.getRegressionData(
            mapDim, numContours=contourCount
        )
        axis.imshow(
            np.ma.filled(face, np.nan),
            cmap="Greys",
            origin="upper",
            interpolation="nearest",
        )
        for levelContours in contours:
            for contour in levelContours:
                axis.plot(
                    contour[:, 1],
                    contour[:, 0],
                    color="#2F75B5",
                    linewidth=0.65,
                )
        axis.set_title(f"Grain {index + 1}: {grain.geomName}")
        axis.set_aspect("equal")
        axis.axis("off")
    for axis in axes[len(motor.grains):]:
        axis.axis("off")
    figure.tight_layout()
    figure.savefig(path, dpi=dpi, transparent=False)


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document, headers, rows, styles):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, header in enumerate(headers):
        set_cell_text(
            table.rows[0].cells[index],
            header,
            bold=True,
            color="FFFFFF",
        )
        set_cell_shading(
            table.rows[0].cells[index], styles["primary_color"]
        )
    for rowIndex, values in enumerate(rows):
        cells = table.add_row().cells
        for column, value in enumerate(values):
            set_cell_text(cells[column], value)
            if rowIndex % 2:
                set_cell_shading(cells[column], styles["light_fill"])
    document.add_paragraph()
    return table


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def add_toc(document):
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Update this field in Word to display the table of contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, placeholder, end))


def configure_document(document, template, metadata):
    styles = template["styles"]
    section = document.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = document.styles["Normal"]
    normal.font.name = styles["font"]
    normal.font.size = Pt(9.5)
    for styleName, size in (
        ("Title", 24),
        ("Subtitle", 13),
        ("Heading 1", 15),
        ("Heading 2", 12),
    ):
        style = document.styles[styleName]
        style.font.name = styles["font"]
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(
            styles["primary_color"]
        )

    header = section.header.paragraphs[0]
    header.text = (
        f"{metadata['document_id']} | Rev. {metadata['revision']} | "
        f"{metadata['motor_name']}"
    )
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Page / Pagina ")
    add_page_field(footer)


def add_cover(document, template, metadata):
    doc = template["document"]
    for _ in range(3):
        document.add_paragraph()
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(doc["title"] + "\n")
    title.add_run(doc["title_pt"])
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(doc["subtitle"] + "\n")
    subtitle.add_run(doc["subtitle_pt"])
    document.add_paragraph()
    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(metadata["motor_name"])
    run.bold = True
    run.font.size = Pt(18)
    document.add_paragraph()
    add_table(
        document,
        ["Field / Campo", "Value / Valor"],
        [
            ("Project / Projeto", metadata["project"]),
            ("Document ID / Documento", metadata["document_id"]),
            ("Revision / Revisao", metadata["revision"]),
            ("Status", metadata["status"]),
            ("Prepared by / Elaborado por", metadata["prepared_by"]),
            ("Organization / Organizacao", metadata["organization"]),
            ("Date / Data", metadata["date"]),
        ],
        template["styles"],
    )
    classification = document.add_paragraph()
    classification.alignment = WD_ALIGN_PARAGRAPH.CENTER
    classification.add_run(
        doc["confidentiality"] + "\n" + doc["confidentiality_pt"]
    ).bold = True
    document.add_page_break()


def add_bilingual_text(document, english, portuguese):
    document.add_paragraph(english)
    paragraph = document.add_paragraph(portuguese)
    for run in paragraph.runs:
        run.italic = True


def metadata_from_config(config):
    supplied = config.get("metadata", {})
    motorName = supplied.get("motor_name") or Path(config["motor"]).stem
    reportDate = supplied.get("date", "auto")
    if reportDate == "auto":
        reportDate = date.today().isoformat()
    return {
        "motor_name": motorName,
        "project": supplied.get("project", "openMotor Engineering Analysis"),
        "prepared_by": supplied.get("prepared_by", ""),
        "organization": supplied.get("organization", ""),
        "document_id": supplied.get("document_id", "OMR-001"),
        "revision": supplied.get("revision", "A"),
        "status": supplied.get("status", "Preliminary"),
        "date": reportDate,
    }


def engineering_checks(motor, result):
    config = motor.config
    checks = [
        (
            "Peak chamber pressure / Pressao maxima",
            result.getMaxPressure(),
            config.getProperty("maxPressure"),
            "<=",
            "Pa",
        ),
        (
            "Peak mass flux / Fluxo de massa maximo",
            result.getPeakMassFlux(),
            config.getProperty("maxMassFlux"),
            "<=",
            "kg/(m^2 s)",
        ),
        (
            "Peak core Mach / Mach maximo no canal",
            result.getPeakMachNumber(),
            config.getProperty("maxMachNumber"),
            "<=",
            "",
        ),
        (
            "Initial port/throat ratio / Razao canal/garganta",
            result.getPortRatio(),
            config.getProperty("minPortThroat"),
            ">=",
            "",
        ),
    ]
    rows = []
    for name, actual, limit, operator, unit in checks:
        passes = (
            actual is not None
            and (
                actual <= limit if operator == "<=" else actual >= limit
            )
        )
        rows.append(
            (
                name,
                format_value(actual, unit),
                f"{operator} {format_value(limit, unit)}",
                "PASS / APROVADO" if passes else "FAIL / REPROVADO",
            )
        )
    return rows


def add_property_table(document, title, properties, template):
    document.add_heading(title, level=2)
    rows = [
        (name, property_value(name, value))
        for name, value in properties.items()
    ]
    add_table(
        document,
        ["Parameter / Parametro", "Value / Valor"],
        rows,
        template["styles"],
    )


def build_document(config, template, motor, result, figures):
    metadata = metadata_from_config(config)
    document = Document()
    configure_document(document, template, metadata)
    add_cover(document, template, metadata)
    sections = template["sections"]
    text = template["text"]

    document.add_heading(sections["revision"], level=1)
    add_table(
        document,
        ["Revision / Revisao", "Date / Data", "Description / Descricao"],
        [
            (
                metadata["revision"],
                metadata["date"],
                "Initial simulation report / Relatorio inicial de simulacao",
            )
        ],
        template["styles"],
    )
    document.add_heading(sections["contents"], level=1)
    add_toc(document)
    document.add_page_break()

    document.add_heading(sections["purpose"], level=1)
    add_bilingual_text(
        document, text["purpose_en"], text["purpose_pt"]
    )
    add_bilingual_text(document, text["method_en"], text["method_pt"])

    document.add_heading(sections["inputs"], level=1)
    add_table(
        document,
        ["Item / Item", "Value / Valor"],
        [
            ("Source motor file / Arquivo fonte", str(config["motor"])),
            ("Grain count / Numero de graos", str(len(motor.grains))),
            (
                "Maximum grain diameter / Diametro maximo",
                format_value(result.getMaxPropellantDiameter() * 1000, "mm"),
            ),
            (
                "Total propellant length / Comprimento total",
                format_value(result.getPropellantLength() * 1000, "mm"),
            ),
        ],
        template["styles"],
    )

    document.add_heading(sections["propellant"], level=1)
    propellant = motor.propellant.getProperties()
    add_property_table(
        document,
        "General properties / Propriedades gerais",
        {key: value for key, value in propellant.items() if key != "tabs"},
        template,
    )
    tabRows = []
    for index, tab in enumerate(propellant["tabs"], start=1):
        tabRows.append(
            (
                index,
                format_value(tab["minPressure"] / 1e6, "MPa"),
                format_value(tab["maxPressure"] / 1e6, "MPa"),
                format_value(tab["a"]),
                format_value(tab["n"]),
                format_value(tab["k"]),
                format_value(tab["t"], "K"),
                format_value(tab["m"], "g/mol"),
            )
        )
    add_table(
        document,
        ["Tab", "P min", "P max", "a", "n", "k", "T", "M"],
        tabRows,
        template["styles"],
    )

    document.add_heading(sections["grains"], level=1)
    grainRows = []
    for index, grain in enumerate(motor.grains, start=1):
        properties = grain.getProperties()
        grainRows.append(
            (
                index,
                grain.geomName,
                property_value("diameter", properties.get("diameter")),
                property_value("length", properties.get("length")),
                "; ".join(
                    f"{key}={property_value(key, value)}"
                    for key, value in properties.items()
                    if key not in ("diameter", "length")
                ),
            )
        )
    add_table(
        document,
        [
            "#",
            "Type / Tipo",
            "Diameter / Diametro",
            "Length / Comprimento",
            "Additional properties / Propriedades adicionais",
        ],
        grainRows,
        template["styles"],
    )

    document.add_heading(sections["nozzle"], level=1)
    add_property_table(
        document,
        "Nozzle geometry / Geometria da tubeira",
        motor.nozzle.getProperties(),
        template,
    )

    document.add_heading(sections["performance"], level=1)
    add_table(
        document,
        ["Metric / Metrica", "Value / Valor"],
        [
            (name, format_value(value, unit))
            for name, value, unit in performance_metrics(result)
        ],
        template["styles"],
    )

    document.add_heading(sections["figures"], level=1)
    document.add_picture(
        str(figures["performance"]), width=Inches(6.4)
    )
    caption = document.add_paragraph(
        "Figure 1. Chamber pressure and thrust histories / "
        "Figura 1. Historicos de pressao da camara e empuxo"
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_picture(
        str(figures["diagnostics"]), width=Inches(6.4)
    )
    caption = document.add_paragraph(
        "Figure 2. Internal ballistic diagnostic histories / "
        "Figura 2. Historicos de diagnostico de balistica interna"
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading(sections["regression"], level=1)
    document.add_picture(
        str(figures["regression"]), width=Inches(6.4)
    )
    caption = document.add_paragraph(
        "Figure 3. Grain cross-sections and regression contours / "
        "Figura 3. Secoes dos graos e contornos de regressao"
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading(sections["compliance"], level=1)
    add_table(
        document,
        [
            "Check / Verificacao",
            "Calculated / Calculado",
            "Limit / Limite",
            "Status",
        ],
        engineering_checks(motor, result),
        template["styles"],
    )
    alertRows = [
        (
            alert.level.name,
            alert.type.name,
            alert.location or "Motor",
            alert.description,
        )
        for alert in result.alerts
    ]
    if not alertRows:
        alertRows = [
            (
                "MESSAGE",
                "VALUE",
                "Motor",
                "No simulation alerts / Nenhum alerta de simulacao",
            )
        ]
    add_table(
        document,
        ["Level / Nivel", "Type / Tipo", "Location / Local", "Description / Descricao"],
        alertRows,
        template["styles"],
    )

    document.add_heading(sections["conclusions"], level=1)
    failedChecks = [
        row for row in engineering_checks(motor, result) if row[-1].startswith("FAIL")
    ]
    if failedChecks:
        add_bilingual_text(
            document,
            (
                "The simulation completed, but one or more configured "
                "engineering limits were not satisfied. Review the compliance "
                "table before proceeding."
            ),
            (
                "A simulacao foi concluida, mas um ou mais limites de "
                "engenharia configurados nao foram atendidos. Revise a tabela "
                "de verificacao antes de prosseguir."
            ),
        )
    else:
        add_bilingual_text(
            document,
            (
                "The simulation completed successfully and the reported "
                "engineering checks satisfy their configured limits."
            ),
            (
                "A simulacao foi concluida com sucesso e as verificacoes de "
                "engenharia apresentadas atendem aos limites configurados."
            ),
        )
    add_bilingual_text(
        document, text["disclaimer_en"], text["disclaimer_pt"]
    )

    document.add_heading(sections["appendix"], level=1)
    add_property_table(
        document,
        "Simulation configuration / Configuracao da simulacao",
        motor.config.getProperties(),
        template,
    )
    return document


def create_figures(config, motor, result, directory):
    settings = config.get("figures", {})
    dpi = int(settings.get("dpi", 180))
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    figures = {
        "performance": directory / "performance.png",
        "diagnostics": directory / "diagnostics.png",
        "regression": directory / "grain-regression.png",
    }
    generate_performance_figure(result, figures["performance"], dpi)
    generate_diagnostics_figure(result, figures["diagnostics"], dpi)
    generate_regression_figure(
        motor,
        figures["regression"],
        dpi,
        int(settings.get("regression_map_dim", 300)),
        int(settings.get("regression_contours", 7)),
    )
    return figures


def generate_report(config, progress=None):
    validate_report_config(config)
    template = load_template(config["template"])
    output = Path(config["output"])
    output.parent.mkdir(parents=True, exist_ok=True)
    if progress:
        progress("Running openMotor simulation", 1, 4)
    motor, result = simulate(config)

    keepFigures = bool(config.get("figures", {}).get("keep_png_files", True))
    if keepFigures:
        figureDirectory = output.parent / f"{output.stem}_figures"
        temporary = None
    else:
        temporary = tempfile.TemporaryDirectory()
        figureDirectory = Path(temporary.name)
    try:
        if progress:
            progress("Generating engineering figures", 2, 4)
        figures = create_figures(
            config, motor, result, figureDirectory
        )
        if progress:
            progress("Filling bilingual document template", 3, 4)
        document = build_document(
            config, template, motor, result, figures
        )
        document.save(output)
    finally:
        if temporary is not None:
            temporary.cleanup()
    if progress:
        progress("Report complete", 4, 4)
    return {
        "output": str(output),
        "figures": (
            str(figureDirectory) if keepFigures else None
        ),
        "designation": result.getDesignation(),
        "impulse_Ns": result.getImpulse(),
        "burn_time_s": result.getBurnTime(),
        "peak_pressure_Pa": result.getMaxPressure(),
        "average_pressure_Pa": result.getAveragePressure(),
        "propellant_mass_kg": result.getPropellantMass(),
    }


def console_progress(status, current, total):
    print(f"[{current}/{total}] {status}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command")
    runParser = subparsers.add_parser("run", help="Generate a report")
    runParser.add_argument("config", help="Path to report YAML")
    guiParser = subparsers.add_parser("gui", help="Open the report GUI")
    guiParser.add_argument("config", nargs="?", help="Optional report YAML")
    args = parser.parse_args()

    if args.command in (None, "gui"):
        from tools.openmotor_report_gui import launch_gui

        launch_gui(getattr(args, "config", None))
        return

    config = load_report_config(args.config)
    result = generate_report(config, console_progress)
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
